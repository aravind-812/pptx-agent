"""Extract text from a PPTX file for eval judges."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt


def extract_pptx_text(pptx_path: str) -> dict:
    """
    Returns:
      {
        "slide_1": {"ShapeName": "text content", ...},
        ...
        "all_text": "flat concat of all text across all slides"
      }
    """
    result: dict = {}
    all_parts: list[str] = []

    try:
        prs = Presentation(pptx_path)
    except Exception as exc:
        return {"error": str(exc), "all_text": ""}

    for i, slide in enumerate(prs.slides, 1):
        slide_key = f"slide_{i}"
        shapes: dict[str, str] = {}
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            lines = []
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if text:
                    lines.append(text)
            if lines:
                text_content = "\n".join(lines)
                shapes[shape.name] = text_content
                all_parts.append(text_content)
        result[slide_key] = shapes

    result["all_text"] = "\n".join(all_parts)
    return result


if __name__ == "__main__":
    import json
    import sys
    if len(sys.argv) < 2:
        print("Usage: python -m eval_agent.extract_pptx_text <pptx_path>")
        sys.exit(1)
    data = extract_pptx_text(sys.argv[1])
    print(f"Slides: {len([k for k in data if k.startswith('slide_')])}")
    print(f"Total text length: {len(data.get('all_text', ''))}")
    print(json.dumps({k: v for k, v in data.items() if k != 'all_text'}, indent=2)[:2000])
